# 生成两个测试用的 xlsx：
#   1) test\schedule_columns.xlsx  规范列名（课程名称/教师/教室/星期/节次/周次）—— 用于“仅解析表格”
#   2) test\schedule_grid.xlsx     近似教务系统网格版式（星期做列、节次做行）—— 用于“AI 解析”
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))


def style_sheet(ws):
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for c in ws[1]:
        c.font = Font(bold=True)


# ---------- 1) 规范列名表 ----------
wb = Workbook()
ws = wb.active
ws.title = "课表"
ws.append(["课程名称", "教师", "教室", "星期", "节次", "周次", "备注"])
rows = [
    ["高等数学(A)I", "封丽", "Z206", "星期二", "1-2节", "2-11周", ""],
    ["高等数学(A)I", "封丽", "Z206", "星期三", "1-2节", "1-16周", ""],
    ["高等数学(A)I", "封丽", "Z206", "星期四", "3-4节", "1-16周", ""],
    ["大学英语I", "祝捷", "611", "星期二", "3-4节", "1-16周", ""],
    ["大学英语I", "祝捷", "Z106", "星期二", "6-7节", "2-10周", ""],
    ["大学英语I", "祝捷", "Z210", "星期三", "3-4节", "3-15周（双周）", "双周上课"],
    ["思想道德与法治", "丁一", "Z212", "星期二", "6-8节", "1-16周", ""],
    ["思想道德与法治", "丁一", "Z204", "星期四", "6-8节", "1-16周", ""],
    ["线性代数与几何(A)", "孟昕娜", "Z202", "星期四", "1-2节", "2-10周", ""],
    ["线性代数与几何(A)", "孟昕娜", "Z202", "星期五", "3-4节", "1-16周", ""],
    ["计算机科学与技术导论", "刘展威", "822", "星期五", "6-8节", "2-12周", ""],
    ["C语言程序设计(A)", "周瑛", "Z202", "星期一", "3-4节", "1-16周", ""],
    ["C语言程序设计(A)", "周瑛", "Z404", "星期五", "1-2节", "12-16周", ""],
    ["形势与政策", "丁一", "待定", "星期六", "1-2节", "1-16周", "集中安排"],
    ["体育I", "赵星", "操场", "星期三", "3-4节", "1-16周", ""],
]
for r in rows:
    ws.append(r)
for col, w in zip("ABCDEFG", [22, 10, 10, 10, 10, 12, 14]):
    ws.column_dimensions[col].width = w
style_sheet(ws)
wb.save(os.path.join(HERE, "schedule_columns.xlsx"))

# ---------- 2) 网格版式（星期为列） ----------
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "学生课表"
ws2.append(["石家庄铁道大学四方学院 2026-2027 学年第 1 学期课程表"])
ws2.append(["班级：方2603-2", "校区：南校区", "专业：计算机科学与技术"])
ws2.append(["节次", "星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"])
grid = [
    ["1-2节", "", "高等数学(A)I\nZ206\n第2-11周", "高等数学(A)I\nZ206\n第1-16周", "线性代数与几何(A)\nZ202\n第2-10周", "", "", ""],
    ["3-4节", "C语言程序设计(A)\nZ202\n第1-16周", "大学英语I\n611\n第1-16周", "体育I\n操场\n第1-16周", "高等数学(A)I\nZ206\n第1-16周", "线性代数与几何(A)\nZ202\n第1-16周", "", ""],
    ["6-8节", "", "思想道德与法治\nZ212\n第1-16周", "", "思想道德与法治\nZ204\n第1-16周", "", "", ""],
]
for r in grid:
    ws2.append(r)
ws2.append(["", "", "大学英语I Z106 第2-10周（6-7节）", "大学英语I Z210 双周（3-4节）", "", "C语言程序设计(A) Z404 第12-16周（1-2节）", "", ""])
for col, w in zip("ABCDEFGH", [10, 20, 22, 22, 22, 22, 12, 12]):
    ws2.column_dimensions[col].width = w
style_sheet(ws2)
wb2.save(os.path.join(HERE, "schedule_grid.xlsx"))

print("OK")
print(os.path.join(HERE, "schedule_columns.xlsx"))
print(os.path.join(HERE, "schedule_grid.xlsx"))
